import base64,io,json,tempfile,unittest,zipfile
from pathlib import Path
from unittest.mock import patch
import app
from core import parse_snapshot_input,collect_snapshot,snapshot_workbook,SnapshotCancelled
A='B000000001';B='B000000002';C='B000000003'

def observation(asin,market='IT',members=None):
 return {'valid':True,'title':'Prodotto '+asin,'asins':members or [A,B], 'attributes':{asin:{'style':'Classico' if asin==A else 'Moderno','color':'Giallo','size':'50 cm'}},'url':f'https://www.amazon.it/dp/{asin}','at':'2026-10-08T12:00:00+02:00'}

class SnapshotInputTests(unittest.TestCase):
 def test_csv_duplicates_and_selected_name(self):
  r=parse_snapshot_input(f'ASIN;Name\n{A.lower()};Uno\n{A};Due\n{B};Tre\n'.encode(),'list.csv')
  self.assertEqual(r['duplicates'],1);self.assertEqual(r['seeds'],[{'asin':A,'name':'Uno'},{'asin':B,'name':'Tre'}])
 def test_invalid_row_stops_whole_import(self):
  with self.assertRaisesRegex(ValueError,'3'):parse_snapshot_input(f'ASIN\n{A}\nBAD'.encode(),'a.csv')
 def test_single_column_without_header(self):self.assertEqual(len(parse_snapshot_input(f'{A}\n{B}'.encode(),'a.csv')['seeds']),2)
 def test_xlsx_input_and_formula_value(self):
  from openpyxl import Workbook
  w=Workbook();w.active.append(['Name','ASIN']);w.active.append(['Uno',A]);b=io.BytesIO();w.save(b);w.close()
  self.assertEqual(parse_snapshot_input(b.getvalue(),'a.xlsx')['seeds'],[{'asin':A,'name':'Uno'}])
 def test_requires_asin_header(self):
  with self.assertRaisesRegex(ValueError,'ASIN'):parse_snapshot_input(b'Product;ID\nTest;B000000001','a.csv')
 def test_duplicate_asin_headers_rejected(self):
  with self.assertRaises(ValueError):parse_snapshot_input(f'ASIN;Child ASIN\n{A};{B}'.encode(),'a.csv')

class CollectionTests(unittest.TestCase):
 def test_expands_children_and_reads_own_attributes(self):
  seen=[]
  def read(a):seen.append(a);return observation(a,members=[A,B] if a==A else [A,B,C])
  r=collect_snapshot(A,'IT',read)
  self.assertEqual(seen,[A,B,C]);self.assertEqual(r['asins'],[A,B,C]);self.assertEqual(r['status'],'Gelesen')
  self.assertEqual(r['variants'][1]['attributes']['style'],'Moderno')
 def test_failure_retained_and_not_called_no_variants(self):
  r=collect_snapshot(A,'IT',lambda a:{'error':'Produkt nicht geladen'})
  self.assertEqual(r['status'],'Unklar');self.assertEqual(r['asins'],[]);self.assertEqual(r['variants'][0]['note'],'Produkt nicht geladen')
 def test_child_failure_and_wrong_language(self):
  r=collect_snapshot(A,'IT',lambda a:observation(a) if a==A else {**observation(a),'language_error':'English'})
  self.assertEqual(r['status'],'Teilweise');self.assertEqual(r['variants'][1]['attributes'],{});self.assertEqual(r['variants'][1]['status'],'Unklar')
 def test_limit_and_cancel(self):
  r=collect_snapshot(A,'IT',lambda a:observation(a,members=[A,B,C]),max_members=1)
  self.assertEqual(r['status'],'Teilweise');self.assertEqual(len(r['variants']),3);self.assertIn('Grenze',r['note'])
  with self.assertRaises(SnapshotCancelled):collect_snapshot(A,'IT',lambda a:observation(a),lambda:True)

class SnapshotStorageTests(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.old=app.DATA;app.DATA=Path(self.temp.name)
  with app.db() as c:
   c.executescript('CREATE TABLE snapshot_jobs(id TEXT PRIMARY KEY,name TEXT,created TEXT,finished TEXT,status TEXT,markets TEXT,error TEXT); CREATE TABLE snapshot_items(job_id TEXT,market TEXT,asin TEXT,name TEXT,position INTEGER,status TEXT,at TEXT,note TEXT,found INTEGER,result TEXT,PRIMARY KEY(job_id,market,asin));CREATE TABLE families(id TEXT PRIMARY KEY,name TEXT,market TEXT,variants TEXT);CREATE TABLE settings(key TEXT PRIMARY KEY,value TEXT);')
   c.execute('INSERT INTO families VALUES(?,?,?,?)',('existing','Soll','IT','{}'))
   c.execute('INSERT INTO settings VALUES(?,?)',('config',json.dumps({'visible':False})))
  app.PROGRESS['running']=False;app.SNAPSHOT_ACTIVE['id']=None;app.SNAPSHOT_STOP.clear()
 def tearDown(self):
  app.DATA=self.old;app.PROGRESS['running']=False;app.SNAPSHOT_ACTIVE['id']=None;app.SNAPSHOT_STOP.clear();self.temp.cleanup()
 def create(self,markets=None):
  with patch.object(app.threading,'Thread'):
   return app.start_snapshot({'filename':'list.csv','data':base64.b64encode(f'ASIN;Name\n{A};=SUM(A1:A2)\n{B};Second'.encode()).decode(),'markets':markets or ['IT','DE'],'name':'Review'})
 def test_full_matrix_cache_market_isolation_and_excel(self):
  ident=self.create();calls=[]
  def scanner(page,market,asin):calls.append((market,asin));return observation(asin,market)
  self.assertTrue(app.snapshot_process(ident,lambda m:m,scanner))
  self.assertEqual(calls,[('IT',A),('IT',B),('DE',A),('DE',B)])
  self.assertEqual(app.families()[0]['name'],'Soll');self.assertEqual(app.snapshot_jobs()[0]['done'],4)
  data,kind=app.snapshot_export(ident)
  self.assertEqual(kind,'application/zip')
  from openpyxl import load_workbook
  with zipfile.ZipFile(io.BytesIO(data)) as z:
   self.assertEqual(set(z.namelist()),{'Varianten-Ist-IT.xlsx','Varianten-Ist-DE.xlsx'})
   for market in ('IT','DE'):
    w=load_workbook(io.BytesIO(z.read('Varianten-Ist-'+market+'.xlsx')))
    self.assertEqual(w.sheetnames,['Variationsgruppen','Ländervergleich','Übersicht','Varianten'])
    self.assertEqual(w['Varianten'].max_row,3) # same family from two input ASINs deduplicates
    self.assertEqual(w['Varianten']['C2'].value,market)
    self.assertEqual(w['Varianten']['B2'].value,A+', '+B)
    self.assertEqual(w['Übersicht']['B6'].value,'=SUM(A1:A2)');self.assertEqual(w['Übersicht']['B6'].data_type,'s')
    self.assertEqual(w['Varianten']['N2'].value,'Offen');self.assertEqual(len(w['Varianten'].data_validations.dataValidation),1)
    self.assertEqual(w['Varianten'].freeze_panes,'A2');self.assertIsNotNone(w['Varianten']['M2'].hyperlink)
    w.close()
 def test_stop_retains_finished_tasks_and_resume_only_open(self):
  ident=self.create(['IT']);calls=[]
  def scanner(page,market,asin):
   calls.append(asin)
   if asin==B:app.SNAPSHOT_STOP.set()
   return observation(asin,members=[asin])
  self.assertFalse(app.snapshot_process(ident,lambda m:m,scanner));self.assertEqual(app.snapshot_jobs()[0]['done'],1)
  app.PROGRESS['running']=False
  with patch.object(app.threading,'Thread'):self.assertEqual(app.start_snapshot({},resume_id=ident),ident)
  calls=[]
  self.assertTrue(app.snapshot_process(ident,lambda m:m,lambda p,m,a:(calls.append(a) or observation(a,members=[a]))))
  self.assertEqual(calls,[B]);self.assertEqual(app.snapshot_jobs()[0]['done'],2)
 def test_pending_items_exported_without_fake_zero(self):
  ident=self.create(['IT']);data,_=app.snapshot_export(ident,'IT')
  from openpyxl import load_workbook
  w=load_workbook(io.BytesIO(data));self.assertEqual(w['Übersicht']['D6'].value,'Ausstehend');self.assertIsNone(w['Übersicht']['E6'].value);w.close()
 def test_shared_lock_and_market_validation(self):
  ident=self.create()
  with self.assertRaisesRegex(ValueError,'laufende'):self.create()
  with self.assertRaises(ValueError):app.snapshot_export(ident,'UK')
  page=app.snapshot_page({'id':ident,'market':'DE','q':B});self.assertEqual(page['count'],1)

class ComparisonExportTests(unittest.TestCase):
 def item(self,market,members,status='Gelesen',name='Gartenlampe'):
  return {'market':market,'asin':A,'name':name,'status':status,'result':{'asins':members,'variants':[{'asin':a,'title':'Lampe '+a,'attributes':{'color':'Gelb' if market=='DE' else 'Giallo'},'status':status,'url':'https://www.amazon.it/dp/'+a} for a in members]},'note':''}
 def test_same_members_ignore_language_and_order(self):
  from core import snapshot_comparison
  r=snapshot_comparison([self.item('DE',[A,B]),self.item('IT',[B,A])],['DE','IT'])['rows'][0]
  self.assertEqual(r['verdict'],'Einheitliche Variation');self.assertIn('DE / IT',r['same'])
 def test_same_count_different_members_and_no_variation(self):
  from core import snapshot_comparison
  r=snapshot_comparison([self.item('DE',[A,B]),self.item('IT',[A,C]),self.item('FR',[A])],['DE','IT','FR'])['rows'][0]
  self.assertEqual(r['verdict'],'Abweichend');self.assertIn('Keine Variation – Einzelprodukt',r['markets']['FR']['label'])
  self.assertIn('Gruppe B',r['markets']['IT']['label'])
 def test_partial_and_missing_never_uniform(self):
  from core import snapshot_comparison
  r=snapshot_comparison([self.item('DE',[A,B]),self.item('IT',[A,B],'Teilweise')],['DE','IT','FR'])['rows'][0]
  self.assertEqual(r['verdict'],'Noch nicht sicher vergleichbar');self.assertEqual(r['unknown'],'IT, FR')
 def test_horizontal_product_cards_and_old_json_records(self):
  from openpyxl import load_workbook
  a=self.item('IT',[A,B],name='=FORMULA()');a['result']=json.dumps(a['result'])
  b={**self.item('IT',[B,A]),'asin':B}
  w=load_workbook(io.BytesIO(snapshot_workbook({'name':'Old saved run','status':'Abgeschlossen','created':'','markets':'["IT"]'},[a,b],'IT')))
  sheet=w['Variationsgruppen'];self.assertEqual(sheet.max_row,4)
  self.assertEqual(sheet['A4'].data_type,'s');self.assertTrue(sheet['G4'].value.startswith('Lampe '));self.assertIn(A,sheet['G4'].value);self.assertIn(B,sheet['H4'].value)
  self.assertIsNotNone(sheet['G4'].hyperlink);self.assertEqual(w.active.title,'Variationsgruppen');w.close()
 def test_singletons_and_no_data(self):
  from core import snapshot_comparison
  r=snapshot_comparison([self.item('DE',[A]),self.item('IT',[A])],['DE','IT'])['rows'][0]
  self.assertEqual(r['verdict'],'Einheitlich ohne Variation')
  self.assertEqual(snapshot_comparison([],['IT'])['rows'],[])

class CountryListTests(SnapshotStorageTests):
 def upload(self,asin,name='Country product'):
  return {'filename':'country.csv','data':base64.b64encode(f'ASIN;Name\n{asin};{name}'.encode()).decode()}
 def test_country_lists_scoped_to_their_market(self):
  with patch.object(app.threading,'Thread'):
   ident=app.start_snapshot({**self.upload(A),'markets':['DE','UK','IE'],'market_lists':{'UK':self.upload(B),'IE':self.upload(C)}})
  _,items=app.snapshot_data(ident)
  self.assertEqual([(i['market'],i['asin']) for i in items],[('DE',A),('UK',B),('IE',C)])
 def test_append_preserves_completed_results_and_only_scans_new(self):
  ident=self.create(['IT','UK'])
  app.snapshot_process(ident,lambda m:m,lambda p,m,a:observation(a,m,members=[a]))
  _,before=app.snapshot_data(ident);app.PROGRESS['running']=False
  self.assertEqual(app.extend_snapshot({'id':ident,'market_lists':{'UK':self.upload(A,'Changed name'),'IE':self.upload(C)}}),1)
  _,after=app.snapshot_data(ident);self.assertEqual(after[:len(before)],before)
  self.assertEqual(app.extend_snapshot({'id':ident,'market_lists':{'IE':self.upload(C)}}),0)
  calls=[];app.snapshot_process(ident,lambda m:m,lambda p,m,a:(calls.append((m,a)) or observation(a,m,members=[a])))
  self.assertEqual(calls,[('IE',C)])
  package,_=app.snapshot_export(ident)
  with zipfile.ZipFile(io.BytesIO(package)) as z:self.assertIn('Varianten-Ist-IE.xlsx',z.namelist())
 def test_append_blocked_during_scan_and_invalid_market(self):
  ident=self.create()
  with self.assertRaisesRegex(ValueError,'laufenden'):app.extend_snapshot({'id':ident,'market_lists':{'UK':self.upload(C)}})
  app.PROGRESS['running']=False
  with self.assertRaises(ValueError):app.extend_snapshot({'id':ident,'market_lists':{'DE':self.upload(C)}})
 def test_ireland_url_and_locale(self):
  from core import market_product_url,market_browser_options
  self.assertIn('www.amazon.ie/dp/',market_product_url('IE',A));self.assertEqual(market_browser_options('IE')['locale'],'en-IE')

class SingletonDisplayTests(unittest.TestCase):
 def test_singleton_with_style_is_not_variation_in_all_excel_sheets(self):
  from core import snapshot_structure,snapshot_comparison
  from openpyxl import load_workbook
  item={'market':'UK','asin':A,'name':'MG 280','status':'Gelesen','result':json.dumps({'asins':[A],'variants':[{'asin':A,'title':'Massage mat','attributes':{'style':'Massage mat'},'status':'Gelesen','url':'https://www.amazon.co.uk/dp/'+A}]}),'note':''}
  label='Keine Variation – Einzelprodukt'
  self.assertEqual(snapshot_structure('Gelesen',1),label)
  self.assertNotEqual(snapshot_structure('Teilweise',1),label)
  self.assertNotEqual(snapshot_structure('Unklar',1),label)
  self.assertNotEqual(snapshot_structure('Gelesen',0),label)
  w=load_workbook(io.BytesIO(snapshot_workbook({'name':'Saved scan','status':'Abgeschlossen','created':''},[item],'UK')))
  self.assertEqual(w['Variationsgruppen']['C4'].value,label)
  self.assertEqual(w['Übersicht']['I6'].value,label)
  self.assertEqual(w['Varianten']['P2'].value,label)
  self.assertIn(label,w['Ländervergleich']['F4'].value)
  self.assertEqual(w['Varianten']['G2'].value,'Massage mat');w.close()

class BaselineTests(SnapshotStorageTests):
 def ready(self):
  ident=self.create(['IT','DE']);app.snapshot_process(ident,lambda m:m,lambda p,m,a:observation(a,m));app.PROGRESS['running']=False
  return ident
 def test_adopt_preserves_snapshots_and_enables_schedule(self):
  ident=self.ready();before=app.snapshot_data(ident)
  plan=app.adopt_snapshot({'id':ident});self.assertEqual(len(plan['groups']),2)
  self.assertEqual(plan['groups'][0]['variants'][A]['attributes'],{})
  app.adopt_snapshot({'id':ident,'apply':True,'token':plan['token'],'daily':True,'time':'10:15'})
  self.assertEqual(app.snapshot_data(ident),before);self.assertEqual(len(app.families()),3)
  self.assertEqual(app.config()['time'],'10:15');self.assertTrue(app.config()['enabled'])
  again=app.adopt_snapshot({'id':ident});self.assertEqual(again['groups'],[]);self.assertEqual(again['unchanged'],2)
 def test_stale_preview_rejected_and_existing_soll_untouched(self):
  ident=self.ready();plan=app.adopt_snapshot({'id':ident})
  with app.db() as c:c.execute('UPDATE families SET name=? WHERE id=?',('Changed','existing'))
  with self.assertRaisesRegex(ValueError,'geändert'):app.adopt_snapshot({'id':ident,'apply':True,'token':plan['token']})
  self.assertEqual(len(app.families()),1)
 def test_conflicts_and_unclear_are_not_adopted(self):
  ident=self.ready()
  with app.db() as c:
   c.execute('UPDATE families SET variants=? WHERE id=?',(json.dumps({A:''}),'existing'))
   c.execute('UPDATE snapshot_items SET status=? WHERE job_id=? AND market=? AND asin=?',('Unklar',ident,'DE',B))
  plan=app.adopt_snapshot({'id':ident});self.assertFalse(plan['groups']);self.assertTrue(plan['skipped'])
 def test_opt_in_attributes_and_split_detection(self):
  from core import evaluate
  ident=self.ready();plan=app.adopt_snapshot({'id':ident,'attributes':True})
  self.assertEqual(plan['groups'][0]['variants'][A]['attributes']['style'],'Classico')
  self.assertEqual(evaluate(plan['groups'][0]['variants'],{A:observation(A,members=[A]),B:observation(B,members=[B])})['status'],'Abweichung')
