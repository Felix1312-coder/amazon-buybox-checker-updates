import io,json
from unittest.mock import patch
from openpyxl import load_workbook
import app
from test_snapshots import SnapshotStorageTests,A,B,observation

class QuickExportTests(SnapshotStorageTests):
 def test_combined_preserves_countries_styles_and_safe_strings(self):
  ident=self.create()
  app.snapshot_process(ident,lambda m:m,lambda page,market,asin:observation(asin,market))
  with patch.object(app,'scan_page',side_effect=AssertionError('Export must not rescan')):
   data,kind=app.snapshot_export(ident,format='xlsx')
  w=load_workbook(io.BytesIO(data))
  self.assertEqual(len(w.sheetnames),7)
  self.assertEqual(w.sheetnames[0],'Ländervergleich')
  for m in ['DE','IT']:
   self.assertEqual(w[m+' Varianten']['C2'].value,m)
   self.assertEqual(w[m+' Übersicht']['B6'].value,'=SUM(A1:A2)')
   self.assertEqual(w[m+' Übersicht']['B6'].data_type,'s')
   self.assertEqual(w[m+' Varianten'].freeze_panes,'A2')
   self.assertEqual(len(w[m+' Varianten'].data_validations.dataValidation),1)
   self.assertIsNotNone(w[m+' Varianten']['M2'].hyperlink)
  w.close()

 def test_manual_text_creates_full_matrix_without_changing_soll(self):
  before=app.families()
  with patch.object(app.threading,'Thread'):
   ident=app.start_snapshot({'text':A+'\n'+B+'\n'+A,'markets':['UK','IE'],'name':'Schnelltest'})
  with app.db() as c:rows=c.execute('SELECT market,asin,status FROM snapshot_items WHERE job_id=?',(ident,)).fetchall()
  self.assertEqual({(r['market'],r['asin']) for r in rows},{(m,a) for m in ['UK','IE'] for a in [A,B]})
  data,_=app.snapshot_export(ident,format='xlsx')
  w=load_workbook(io.BytesIO(data))
  self.assertEqual(w['UK Übersicht']['D6'].value,'Ausstehend')
  self.assertIsNone(w['UK Übersicht']['E6'].value)
  self.assertEqual(app.families(),before)
  w.close()
